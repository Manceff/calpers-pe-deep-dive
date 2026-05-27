"""
Ingest calpers_pe_funds.xlsx → aim/data/master_long.parquet

Lit les 20 onglets (un par snapshot), normalise les types, parse les footnotes
(N/M, suffixes "1" non-breaking-space), écrit un parquet long format avec
as_of_date dérivé du nom de l'onglet.
"""
from __future__ import annotations

import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "calpers_pe_funds.xlsx"
OUT = ROOT / "aim" / "data" / "master_long.parquet"


# --- Helpers de parsing ---------------------------------------------------

SHEET_DATE_RE = re.compile(
    r"(January|February|March|April|May|June|July|August|"
    r"September|October|November|December)\s+(\d{1,2}),\s+(\d{4})"
)
MONTHS = {
    "January": 1, "February": 2, "March": 3, "April": 4, "May": 5, "June": 6,
    "July": 7, "August": 8, "September": 9, "October": 10, "November": 11, "December": 12,
}


def sheet_to_date(sheet_name: str) -> date:
    """Convertit 'December 31, 2014,' ou 'September 30, 2025' en date."""
    m = SHEET_DATE_RE.search(sheet_name)
    if not m:
        raise ValueError(f"Impossible d'extraire une date depuis: {sheet_name!r}")
    month, day, year = m.group(1), int(m.group(2)), int(m.group(3))
    return date(year, MONTHS[month], day)


def clean_str(x) -> str | None:
    if x is None:
        return None
    s = unicodedata.normalize("NFKC", str(x))
    s = s.replace("\xa0", " ").strip()
    return s if s else None


def parse_money(x) -> float | None:
    s = clean_str(x)
    if s is None or "N/M" in s:
        return None
    # Les cellules money du source CalPERS n'ont pas de footnote — juste $ et ,
    s = s.replace("$", "").replace(",", "").strip()
    if not s or s == "-":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_pct(x) -> float | None:
    s = clean_str(x)
    if s is None or "N/M" in s:
        return None
    # Drop trailing footnote: "5.7%1" or "5.7% 1"
    s = re.sub(r"%\s*\d*$", "", s).strip()
    try:
        return float(s)
    except ValueError:
        return None


def parse_mult(x) -> float | None:
    s = clean_str(x)
    if s is None or "N/M" in s:
        return None
    # Drop trailing footnote: "1.4x1" or "1.4x 1"
    s = re.sub(r"x\s*\d*$", "", s).strip()
    try:
        return float(s)
    except ValueError:
        return None


def make_fund_key(name: str) -> str:
    """Slug stable depuis fund_name (lowercase, ascii, sans ponctuation)."""
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    s = s.lower()
    # Retire les contenus parenthétiques connus qui sont des qualificatifs
    # CalPERS-only (ex: "Fund 2 (CalPERS), LLC" → "Fund 2 LLC")
    s = re.sub(r"\(\s*calpers\s*\)", " ", s)
    s = re.sub(r"[^\w\s]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    # Drop common legal suffixes pour stabiliser la clé
    s = re.sub(
        r"\b(l\s*p|llc|lp|inc|ltd|limited|partnership|partners|l\s*l\s*c)\b\.?",
        "",
        s,
    )
    s = re.sub(r"\s+", " ", s).strip()
    return s


def harmonize_vintage_year(master: pd.DataFrame) -> pd.DataFrame:
    """Aligne vintage_year sur la valeur modale par fund_key.

    CalPERS publie parfois des restatements d'année de vintage entre snapshots
    (ex: California Emerging Ventures II reporté 2000 en 2014 puis 1999 ensuite,
    Siris Partners III reporté 2014 en 2016-06 puis 2015 ensuite). On retient
    la valeur **la plus fréquente** par fund_key pour garantir la cohérence des
    agrégats par millésime. La valeur initiale erronée disparait au profit de la
    correction officielle CalPERS ultérieure.
    """
    counts = (master.dropna(subset=["vintage_year"])
              .groupby(["fund_key", "vintage_year"]).size().reset_index(name="n"))
    modal = (counts.sort_values(["fund_key", "n", "vintage_year"],
                                 ascending=[True, False, True])
                   .groupby("fund_key").head(1)
                   .set_index("fund_key")["vintage_year"])
    master = master.copy()
    master["vintage_year"] = master["fund_key"].map(modal).astype("Int64")
    return master


def merge_footnote_variants(master: pd.DataFrame) -> pd.DataFrame:
    """Fusionne UNIQUEMENT les vrais footnotes PDF, pas les numéros de série.

    Règle A (série numérotée) : fusionne '<base> N' → '<base>' si :
      - N ∈ {1, 2}
      - le DERNIER token de '<base>' est lui-même un numéro de série (roman ou
        arabe) — donc '<base>' a déjà sa numérotation, ce qui prouve que le
        chiffre trailing est un footnote (et pas une série).
        Ex: 'permira iv 2' → 'permira iv' (le 2 est un footnote, IV est la série).

    Règle B (footnote collé) : fusionne '<base> N' → '<base>' si :
      - N ∈ {1, 2}
      - le fund_name source porte le digit **collé** au texte (sans espace),
        donc un footnote PDF indiscutable (ex: 'Generation Capital Partners, L.P.1')
      - et '<base>' existe ailleurs dans le dataset.

    Évite le bug où '57 stars global opportunities fund 2' (= Fund II) était
    fusionné à tort avec '57 stars global opportunities fund' : ce cas a un
    espace avant le digit dans le nom source ('Fund 2'), donc règle B ne match
    pas, et règle A non plus ('fund' n'est pas roman/arabic).
    """
    ROMAN_OR_ARABIC = re.compile(r"^([ivxlcdm]+|[ivxlcdm]+[\-]?[a-z]|\d+)$",
                                  re.IGNORECASE)
    keys = set(master["fund_key"].unique())

    # Pour règle B : map fund_key → True si AU MOINS un fund_name source
    # a le digit final collé (pas d'espace avant le digit).
    tight_digit_re = re.compile(r"\S\d$")  # caractère non-espace immédiatement avant le digit final
    has_tight_footnote: dict[str, bool] = {}
    for fk, fn in zip(master["fund_key"], master["fund_name"]):
        if fk in has_tight_footnote:
            continue
        if fn and tight_digit_re.search(fn.rstrip()):
            has_tight_footnote[fk] = True

    remap: dict[str, str] = {}
    for k in keys:
        m = re.match(r"^(.+?)\s(\d)$", k)
        if not m:
            continue
        base, digit = m.group(1), m.group(2)
        if digit not in {"1", "2"}:
            continue
        if base not in keys:
            continue  # pas de base co-existante → rien à fusionner
        last_tok = base.split()[-1] if base else ""
        # Règle A : la base se termine par un numéro de série explicite
        if ROMAN_OR_ARABIC.match(last_tok):
            remap[k] = base
            continue
        # Règle B : le fund_name source porte le digit collé (footnote indiscutable)
        if has_tight_footnote.get(k, False):
            remap[k] = base
    if remap:
        master = master.copy()
        master["fund_key"] = master["fund_key"].replace(remap)
    return master


# --- Ingestion -----------------------------------------------------------


def ingest(src: Path = SRC, out: Path = OUT) -> pd.DataFrame:
    if not src.exists():
        raise FileNotFoundError(f"Source introuvable: {src}")

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    frames = []

    for sheet in wb.sheetnames:
        as_of = sheet_to_date(sheet)
        ws = wb[sheet]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue  # header
            if row[0] is None:
                continue
            rows.append(row)
        if not rows:
            continue

        df = pd.DataFrame(
            rows,
            columns=[
                "fund_name",
                "vintage_year",
                "capital_committed",
                "cash_in",
                "cash_out",
                "cash_out_rv",
                "net_irr",
                "investment_multiple",
            ],
        )
        df["fund_name"] = df["fund_name"].map(clean_str)
        df = df.dropna(subset=["fund_name"])
        df["vintage_year"] = pd.to_numeric(df["vintage_year"], errors="coerce").astype("Int64")
        for col in ("capital_committed", "cash_in", "cash_out", "cash_out_rv"):
            df[col] = df[col].map(parse_money)
        df["net_irr"] = df["net_irr"].map(parse_pct)
        df["investment_multiple"] = df["investment_multiple"].map(parse_mult)
        df["as_of_date"] = pd.to_datetime(as_of)
        df["fund_key"] = df["fund_name"].map(make_fund_key)
        frames.append(df)

    master = pd.concat(frames, ignore_index=True)
    master = merge_footnote_variants(master)
    master = harmonize_vintage_year(master)
    master = master[
        [
            "as_of_date",
            "fund_key",
            "fund_name",
            "vintage_year",
            "capital_committed",
            "cash_in",
            "cash_out",
            "cash_out_rv",
            "net_irr",
            "investment_multiple",
        ]
    ].sort_values(["as_of_date", "fund_key"]).reset_index(drop=True)

    out.parent.mkdir(parents=True, exist_ok=True)
    master.to_parquet(out, index=False)
    return master


def main() -> int:
    print(f"Reading {SRC} …", file=sys.stderr)
    df = ingest()
    print(f"Wrote {OUT} — {len(df):,} rows, {df['fund_key'].nunique()} fund_keys, "
          f"{df['as_of_date'].nunique()} snapshots", file=sys.stderr)
    # Summary par snapshot
    by_snap = df.groupby("as_of_date").size()
    for d, n in by_snap.items():
        print(f"  {d.date()}  n={n}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
